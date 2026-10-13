"""Generate academic, submission-ready PowerPoint presentation for review session."""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

WORKSPACE = Path(__file__).resolve().parent.parent
PPTX_OUTPUT = WORKSPACE / "docs" / "Structural_Damage_Recognition_Presentation.pptx"
PLOTS_DIR = WORKSPACE / "outputs" / "plots"

# Academic Color Palette
COLOR_NAVY = RGBColor(26, 54, 93)      # #1A365D Primary Header
COLOR_BLUE = RGBColor(43, 108, 176)    # #2B6CB0 Secondary Accent
COLOR_DARK = RGBColor(45, 55, 72)      # #2D3748 Body Text
COLOR_LIGHT_BG = RGBColor(247, 250, 252) # #F7FAFC Card Background
COLOR_MUTED = RGBColor(113, 128, 150)  # #718096 Subtitles
COLOR_CARD_BORDER = RGBColor(203, 213, 224)


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base_slide(title_text: str, category_text: str = "UE24CS352A — MACHINE LEARNING MINI-PROJECT"):
        slide = prs.slides.add_slide(blank_layout)

        # Top Category Accent
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_BLUE

        # Slide Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

        return slide

    def add_card(slide, left, top, width, height, title="", subtitle=""):
        # Shape background
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_LIGHT_BG
        shape.line.color.rgb = COLOR_CARD_BORDER
        shape.line.width = Pt(1)

        if title or subtitle:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.8))
            tf = tb.text_frame
            tf.word_wrap = True
            if title:
                p = tf.paragraphs[0]
                p.text = title
                p.font.name = "Arial"
                p.font.size = Pt(15)
                p.font.bold = True
                p.font.color.rgb = COLOR_NAVY
            if subtitle:
                p2 = tf.add_paragraph()
                p2.text = subtitle
                p2.font.name = "Arial"
                p2.font.size = Pt(11)
                p2.font.color.rgb = COLOR_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ─────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(blank_layout)

    # Decorative header bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.4))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_BLUE
    top_bar.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_sub = tf1.paragraphs[0]
    p_sub.text = "UE24CS352A — MACHINE LEARNING MINI-PROJECT (SEC-C, TEAM C-22, PROJECT 31)"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(13)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_BLUE

    p_main = tf1.add_paragraph()
    p_main.text = "Structural Damage Recognition Using Transfer Learning"
    p_main.font.name = "Arial"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_NAVY
    p_main.space_before = Pt(8)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Automated Binary Damage State Classification on PEER Hub ImageNet (Phi-Net Task 2)"
    p_desc.font.name = "Arial"
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = COLOR_DARK
    p_desc.space_before = Pt(8)

    # Team Box
    team_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.3), Inches(2.0))
    tf_team = team_box.text_frame
    p_t1 = tf_team.paragraphs[0]
    p_t1.text = "PROJECT METADATA & TEAM:"
    p_t1.font.bold = True
    p_t1.font.size = Pt(13)
    p_t1.font.color.rgb = COLOR_NAVY

    p_t2 = tf_team.add_paragraph()
    p_t2.text = "• Student 1: MIRUNJAI SURESH KUMAR (SRN: PES1UG24AM161)"
    p_t2.font.size = Pt(13)
    p_t2.font.color.rgb = COLOR_DARK
    p_t2.space_before = Pt(4)

    p_t3 = tf_team.add_paragraph()
    p_t3.text = "• Student 2: MOHAMMED SINAN M T (SRN: PES1UG24AM165)"
    p_t3.font.size = Pt(13)
    p_t3.font.color.rgb = COLOR_DARK

    p_t4 = tf_team.add_paragraph()
    p_t4.text = "• Course: UE24CS352A | Section: Sec-C | Team ID: C-22 (Team 22) | Project No: 31 | Date: October 2026"
    p_t4.font.size = Pt(12)
    p_t4.font.color.rgb = COLOR_MUTED
    p_t4.space_before = Pt(8)

    notes1 = s1.notes_slide.notes_text_frame
    notes1.text = (
        "SPEAKER SCRIPT:\n"
        "Good morning faculty and evaluation panel. Today we present our Machine Learning mini-project: "
        "'Structural Damage Recognition Using Transfer Learning' on the benchmark PEER Hub ImageNet dataset. "
        "Our team (Section C, Team C-22) consists of Mirunjai Suresh Kumar and Mohammed Sinan M T. Over the next "
        "few minutes, we will present the disaster management motivation, memory-efficient data pipeline, verified "
        "experimental results on the official 1,460 test images, and demonstrate our live interactive application."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement & Motivation
    # ─────────────────────────────────────────────────────────────
    s2 = add_base_slide("Problem Statement & Disaster Management Motivation")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "The Real-World Challenge")
    tb2_left = s2.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(5.2), Inches(4.0))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True
    p = tf2_l.paragraphs[0]
    p.text = "• Post-Disaster Bottleneck: Following major earthquakes or structural failures, thousands of buildings require immediate screening for occupancy safety."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p2 = tf2_l.add_paragraph()
    p2.text = "• Critical Personnel Shortage: On-site inspection requires licensed civil and structural engineers. Conducting exhaustive visual audits takes weeks to months."
    p2.font.size = Pt(13)
    p2.space_before = Pt(10)
    p3 = tf2_l.add_paragraph()
    p3.text = "• Subjectivity & Hazards: Manual visual triage under emergency conditions is inherently subjective and exposes inspection crews to hazardous aftershocks and collapsing masonry."
    p3.font.size = Pt(13)
    p3.space_before = Pt(10)

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Engineering Objectives")
    tb2_right = s2.shapes.add_textbox(Inches(7.0), Inches(2.5), Inches(5.3), Inches(4.0))
    tf2_r = tb2_right.text_frame
    tf2_r.word_wrap = True
    p = tf2_r.paragraphs[0]
    p.text = "1. Binary Classification: Formulate a computer vision model to automatically categorize structural elements into Damaged (0) vs. Undamaged (1)."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p2 = tf2_r.add_paragraph()
    p2.text = "2. Resource Efficiency: Design a memory-mapped, feature-cached data pipeline that trains efficiently within consumer hardware limits (Intel Core i3, 8 GB RAM) without out-of-memory crashes."
    p2.font.size = Pt(13)
    p2.space_before = Pt(10)
    p3 = tf2_r.add_paragraph()
    p3.text = "3. Rigorous Evaluation: Benchmark strictly on the official held-out test set (1,460 samples) with comprehensive precision, recall, and ROC-AUC metrics."
    p3.font.size = Pt(13)
    p3.space_before = Pt(10)
    p4 = tf2_r.add_paragraph()
    p4.text = "4. Interactive Prototype: Deploy a working Streamlit application for single-photograph inference."
    p4.font.size = Pt(13)
    p4.space_before = Pt(10)

    notes2 = s2.notes_slide.notes_text_frame
    notes2.text = (
        "SPEAKER SCRIPT:\n"
        "Post-earthquake structural triage is a matter of life safety. When an earthquake strikes an urban area, "
        "inspecting thousands of structures manually is slow and perilous. Our objective was to develop an assistive "
        "machine learning tool that rapid-triages photographs of columns, beams, and facades into Damaged or Undamaged states. "
        "Crucially, we set strict software engineering goals: our solution had to execute reliably on consumer laptops "
        "using memory-mapped arrays, achieve high recall on the damaged class, and integrate into an intuitive web interface."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 3: Dataset Summary (Phi-Net Task 2)
    # ─────────────────────────────────────────────────────────────
    s3 = add_base_slide("Dataset Overview: PEER Hub ImageNet (Phi-Net)")
    add_card(s3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.3), "Benchmark Dataset Properties")
    tb3_top = s3.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.3), Inches(1.4))
    tf3_t = tb3_top.text_frame
    tf3_t.word_wrap = True
    p = tf3_t.paragraphs[0]
    p.text = "• Source: PEER Hub ImageNet (Phi-Net) — Task 2: Damage State (Gao & Mosalam, UC Berkeley, 2018; 2020)."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p2 = tf3_t.add_paragraph()
    p2.text = "• Structural Focus: Standardized 224x224x3 photographs of concrete columns, beams, shear walls, and masonry facades."
    p2.font.size = Pt(13)
    p3 = tf3_t.add_paragraph()
    p3.text = "• Label Convention: Class 0 = Damaged (D), Class 1 = Undamaged (UD). Two-column one-hot float32 vectors."
    p3.font.size = Pt(13)

    add_card(s3, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.6), "Training Set (11,811 Samples)")
    tb3_tr = s3.shapes.add_textbox(Inches(1.0), Inches(4.8), Inches(5.2), Inches(1.8))
    tf3_tr = tb3_tr.text_frame
    tf3_tr.word_wrap = True
    p = tf3_tr.paragraphs[0]
    p.text = "• Array Size on Disk: ~6.62 GiB (task2_X_train.npy)"
    p.font.size = Pt(13)
    p2 = tf3_tr.add_paragraph()
    p2.text = "• Damaged (Class 0): 6,282 samples (53.19%)"
    p2.font.size = Pt(13)
    p3 = tf3_tr.add_paragraph()
    p3.text = "• Undamaged (Class 1): 5,529 samples (46.81%)"
    p3.font.size = Pt(13)
    p4 = tf3_tr.add_paragraph()
    p4.text = "• Split: 80% Train (9,448) / 20% Val (2,363) stratified."
    p4.font.size = Pt(13)

    add_card(s3, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.6), "Official Test Set (1,460 Samples)")
    tb3_te = s3.shapes.add_textbox(Inches(7.0), Inches(4.8), Inches(5.3), Inches(1.8))
    tf3_te = tb3_te.text_frame
    tf3_te.word_wrap = True
    p = tf3_te.paragraphs[0]
    p.text = "• Array Size on Disk: ~838.36 MiB (task2_X_test.npy)"
    p.font.size = Pt(13)
    p2 = tf3_te.add_paragraph()
    p2.text = "• Damaged (Class 0): 745 samples (51.03%)"
    p2.font.size = Pt(13)
    p3 = tf3_te.add_paragraph()
    p3.text = "• Undamaged (Class 1): 715 samples (48.97%)"
    p3.font.size = Pt(13)
    p4 = tf3_te.add_paragraph()
    p4.text = "• Protocol: Held strictly out of training and hyperparameter tuning."
    p4.font.size = Pt(13)

    notes3 = s3.notes_slide.notes_text_frame
    notes3.text = (
        "SPEAKER SCRIPT:\n"
        "Here is the dataset breakdown. We utilize Task 2 of the UC Berkeley PEER Hub ImageNet benchmark. "
        "The training set contains 11,811 images occupying 6.6 gigabytes on disk, with a nearly balanced class ratio: "
        "53% damaged and 47% undamaged. The official test set contains 1,460 held-out images. "
        "An important empirical finding during our audit was that the labels in both arrays were stored in contiguous blocks: "
        "all class 0 first, followed by all class 1. Therefore, stratified shuffling was mandatory to prevent single-class batches."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 4: Preprocessing & Hardware Optimization
    # ─────────────────────────────────────────────────────────────
    s4 = add_base_slide("Data Preprocessing & Memory-Efficient Engineering")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Caffe BGR Restoration vs. Inference")
    tb4_l = s4.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.2))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True
    p = tf4_l.paragraphs[0]
    p.text = "• Dataset Storage Quirk: The raw .npy arrays stored images with Caffe-style BGR ImageNet mean subtraction: [103.939, 116.779, 123.68]. Value range was [-123.68, 151.06]."
    p.font.size = Pt(12)
    p2 = tf4_l.add_paragraph()
    p2.text = "• Training Transformation: Added back channel means, clipped to [0, 255], converted BGR -> RGB, and scaled to [-1, 1] via (x / 127.5) - 1.0."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf4_l.add_paragraph()
    p3.text = "• Inference Safety: Uploaded user photographs (JPG/PNG) are ALREADY standard RGB [0, 255]. Our preprocessing module specifically avoids BGR mean addition for uploads, performing Lanczos resize to 224x224 and [-1, 1] scaling directly."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)

    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Resumable Disk-Backed Caching")
    tb4_r = s4.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.2))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True
    p = tf4_r.paragraphs[0]
    p.text = "• The 8 GB RAM Limit: Loading the 6.6 GiB training array plus model weights easily triggers Windows out-of-memory aborts."
    p.font.size = Pt(12)
    p2 = tf4_r.add_paragraph()
    p2.text = "• Memory-Mapped Reading: Arrays are opened with np.load(..., mmap_mode='r'), reading only the current mini-batch (size 8)."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf4_r.add_paragraph()
    p3.text = "• Disk-Backed Output Allocation: Extracted 1,280-dim feature vectors are written directly to disk sectors via np.lib.format.open_memmap (~60 MB cache file)."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)
    p4 = tf4_r.add_paragraph()
    p4.text = "• Resumption Checkpoint: An atomic .meta.json sidecar tracks completed batches. If extraction is interrupted, it resumes from the exact last batch."
    p4.font.size = Pt(12)
    p4.space_before = Pt(8)

    notes4 = s4.notes_slide.notes_text_frame
    notes4.text = (
        "SPEAKER SCRIPT:\n"
        "Turning to preprocessing and systems engineering. The dataset arrived preprocessed with Caffe BGR mean subtraction, "
        "which we inverted back to standard RGB before mapping into MobileNetV2's [-1, 1] input range. "
        "Critically, uploaded JPG photographs are already standard RGB, so our app cleanly separates dataset array restoration "
        "from inference upload preprocessing. "
        "To run on an 8 GB laptop, we engineered a resumable disk-backed caching system using NumPy memory maps. "
        "This kept active RAM usage strictly below 1.5 GB and allowed resuming extraction without re-processing earlier images."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 5: Model Architecture & Transfer Learning
    # ─────────────────────────────────────────────────────────────
    s5 = add_base_slide("Model Architecture & Transfer Learning Strategy")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), "MobileNetV2 Feature Extractor + Lightweight Classification Head")

    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(11.3), Inches(4.2))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    p = tf5.paragraphs[0]
    p.text = "1. Pretrained Backbone: MobileNetV2 (Sandler et al., 2018) instantiated without top head (include_top=False). Contains 2,257,984 frozen parameters in depthwise separable convolution blocks."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK

    p2 = tf5.add_paragraph()
    p2.text = "2. Frozen Weights: All convolutional layers are frozen (trainable = False). Rich low-level and mid-level edge, texture, and crack feature representations from ImageNet are preserved without gradient computation."
    p2.font.size = Pt(13)
    p2.space_before = Pt(10)

    p3 = tf5.add_paragraph()
    p3.text = "3. Dimensionality Reduction: GlobalAveragePooling2D condenses spatial feature maps into a compact 1,280-dimensional feature vector per image."
    p3.font.size = Pt(13)
    p3.space_before = Pt(10)

    p4 = tf5.add_paragraph()
    p4.text = "4. Dense Classification Head (82,114 Trainable Parameters):"
    p4.font.size = Pt(13)
    p4.space_before = Pt(10)

    p4_sub = tf5.add_paragraph()
    p4_sub.text = "   • Dropout Layer (rate = 0.4) — Prevents co-adaptation and regularizes extracted features.\n" \
                  "   • Dense Layer (64 units, ReLU) — Learns non-linear combinations of structural features (81,984 params).\n" \
                  "   • Output Dense Layer (2 units, Softmax) — Produces class prediction scores [Damaged, Undamaged] (130 params).\n" \
                  "   • Total Model Parameters: 2,340,098 (82,114 trainable, 2,257,984 frozen non-trainable)."
    p4_sub.font.size = Pt(12)
    p4_sub.font.color.rgb = COLOR_MUTED
    p4_sub.space_before = Pt(4)

    p5 = tf5.add_paragraph()
    p5.text = "5. Resource Efficiency: Feature caching on disk allows training the classification head on CPU in ~20-30 seconds without out-of-memory errors."
    p5.font.size = Pt(13)
    p5.font.color.rgb = COLOR_BLUE
    p5.space_before = Pt(10)

    notes5 = s5.notes_slide.notes_text_frame
    notes5.text = (
        "SPEAKER SCRIPT:\n"
        "Here is our verified model architecture. We adopted transfer learning using MobileNetV2 with include_top=False. "
        "The frozen convolutional backbone contains exactly 2,257,984 parameters. "
        "Global average pooling condenses each image into a 1,280-dimensional feature representation. "
        "Our trainable head consists of a 40% dropout layer, 64 hidden ReLU units, and a 2-class softmax output, "
        "totaling 82,114 trainable parameters and 2,340,098 total parameters. "
        "Training this head on the cached features takes under 30 seconds on a standard dual-core laptop CPU."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 6: Training Convergence & Loss Curves
    # ─────────────────────────────────────────────────────────────
    s6 = add_base_slide("Training Convergence & Learning Curves")
    curves_file = PLOTS_DIR / "training_curves.png"
    if curves_file.exists():
        s6.shapes.add_picture(str(curves_file), Inches(0.8), Inches(1.6), width=Inches(6.8))

    add_card(s6, Inches(7.8), Inches(1.6), Inches(4.7), Inches(5.2), "Training Dynamics Analysis")
    tb6 = s6.shapes.add_textbox(Inches(8.0), Inches(2.4), Inches(4.3), Inches(4.2))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "• Convergence: 19 epochs total. EarlyStopping restored best model weights at epoch 14."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK

    p2 = tf6.add_paragraph()
    p2.text = "• Loss Trajectory: Training loss decreased monotonically from 0.52 to 0.22. Validation loss stabilized at 0.34."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)

    p3 = tf6.add_paragraph()
    p3.text = "• Generalization Gap: Validation accuracy stabilized at 85.7% against 91.0% training accuracy. The small 5.3% gap demonstrates healthy generalization without severe overfitting."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)

    p4 = tf6.add_paragraph()
    p4.text = "• Learning Rate Scheduling: ReduceLROnPlateau halved learning rate from 0.001 to 0.0005, fine-tuning convergence around the loss minimum."
    p4.font.size = Pt(12)
    p4.space_before = Pt(8)

    notes6 = s6.notes_slide.notes_text_frame
    notes6.text = (
        "SPEAKER SCRIPT:\n"
        "Looking at our training curves: on the left is categorical cross-entropy loss, and on the right is accuracy. "
        "Training progressed smoothly over 19 epochs with EarlyStopping restoring the optimal weights at epoch 14. "
        "The validation accuracy stabilized at 85.7% compared to 91% on training data. "
        "This tight 5.3% generalization gap confirms that our 0.4 dropout rate provided effective regularization."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 7: Evaluation Results on Official Test Set
    # ─────────────────────────────────────────────────────────────
    s7 = add_base_slide("Official Benchmark Test Results (1,460 Samples)")
    cm_file = PLOTS_DIR / "confusion_matrix.png"
    if cm_file.exists():
        s7.shapes.add_picture(str(cm_file), Inches(0.8), Inches(1.6), width=Inches(5.0))

    add_card(s7, Inches(6.0), Inches(1.6), Inches(6.5), Inches(5.2), "Key Performance Indicators")
    tb7 = s7.shapes.add_textbox(Inches(6.2), Inches(2.3), Inches(6.1), Inches(4.3))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "• Official Test Accuracy: 84.25% (1,230 / 1,460 correct)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

    p2 = tf7.add_paragraph()
    p2.text = "• ROC-AUC (Damaged Target): 0.9289 (Excellent discriminative power)"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_BLUE
    p2.space_before = Pt(6)

    p3 = tf7.add_paragraph()
    p3.text = "• Damaged Class (Class 0): Precision = 83.7% | Recall = 85.9% | F1 = 0.85"
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)

    p4 = tf7.add_paragraph()
    p4.text = "• Undamaged Class (Class 1): Precision = 84.9% | Recall = 82.5% | F1 = 0.84"
    p4.font.size = Pt(12)
    p4.space_before = Pt(4)

    p5 = tf7.add_paragraph()
    p5.text = "• Confusion Matrix Breakdown:\n" \
             "   - True Positives (Damaged detected): 640\n" \
             "   - True Negatives (Undamaged detected): 590\n" \
             "   - False Positives (False alarms): 125 (17.5%)\n" \
             "   - False Negatives (Missed damage): 105 (14.1%)"
    p5.font.size = Pt(11)
    p5.font.color.rgb = COLOR_DARK
    p5.space_before = Pt(8)

    notes7 = s7.notes_slide.notes_text_frame
    notes7.text = (
        "SPEAKER SCRIPT:\n"
        "These are our verified results on the official held-out test set of 1,460 images. "
        "The model achieved 84.25% test accuracy and an ROC-AUC of 0.9289. "
        "Notice that our recall for the Damaged class is 85.91%, meaning we correctly catch 640 out of 745 damaged structures. "
        "In post-disaster triage, recall on the damaged class is the most critical metric because a false negative—clearing a damaged building—"
        "is far more dangerous than inspecting an undamaged building that triggered a false alarm."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 8: Error Analysis & Misclassifications
    # ─────────────────────────────────────────────────────────────
    s8 = add_base_slide("Diagnostic Error Analysis")
    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "False Negatives: 105 Images (14.1% Miss Rate)", "Damaged structures misclassified as Undamaged")
    tb8_l = s8.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.2), Inches(4.0))
    tf8_l = tb8_l.text_frame
    tf8_l.word_wrap = True
    p = tf8_l.paragraphs[0]
    p.text = "• Hairline Cracking: Micro-fissures narrower than 2-3 pixels get smoothed out by the 224x224 downsampling."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf8_l.add_paragraph()
    p2.text = "• Lighting & Low Contrast: Severe direct sunlight or heavy shadows obscure surface fractures against gray concrete."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf8_l.add_paragraph()
    p3.text = "• Peripheral Defects: Damage located on the extreme image borders where global average pooling assigns lower spatial saliency."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)

    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "False Positives: 125 Images (17.5% Alarm Rate)", "Undamaged structures misclassified as Damaged")
    tb8_r = s8.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.0))
    tf8_r = tb8_r.text_frame
    tf8_r.word_wrap = True
    p = tf8_r.paragraphs[0]
    p.text = "• Architectural Seams & Mortar: Deep mortar lines in brickwork or textured rustic finishes mimic fracture patterns."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf8_r.add_paragraph()
    p2.text = "• Surface Clutter & Occlusions: Power lines, conduit pipes, drainage stains, or tree shadows cast linear fissure-like silhouettes."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf8_r.add_paragraph()
    p3.text = "• Perspective Distortion: Extreme oblique camera angles distort straight architectural joints into apparent shear deformities."
    p4 = tf8_r.add_paragraph()
    p4.text = "• Diagnostic Scope: These represent plausible visual failure modes identified during diagnostic review rather than on-site physical forensics."
    p4.font.size = Pt(11)
    p4.font.italic = True
    p4.font.color.rgb = COLOR_MUTED
    p4.space_before = Pt(8)

    notes8 = s8.notes_slide.notes_text_frame
    notes8.text = (
        "SPEAKER SCRIPT:\n"
        "Why does an 84% accurate model make mistakes? Our error analysis uncovered two plausible visual failure modes. "
        "The 105 false negatives plausibly occur where fine hairline cracks get smoothed away by 224-pixel downsampling or where harsh shadows conceal defects. "
        "Conversely, the 125 false alarms plausibly occur when high-contrast mortar lines, exposed conduit, or tree branch shadows "
        "mimic crack geometry. Note that these are visual hypotheses from model diagnostics rather than on-site physical forensics."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 9: Interactive Application Demo
    # ─────────────────────────────────────────────────────────────
    s9 = add_base_slide("Interactive Streamlit Application Demo")
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Streamlit Architecture (app.py)")
    tb9_l = s9.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.2))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    p = tf9_l.paragraphs[0]
    p.text = "• Unified Inference Engine: Loads the unified end-to-end model (MobileNetV2 feature extractor + classification head) using @st.cache_resource."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf9_l.add_paragraph()
    p2.text = "• Retracing Prevention: Uses model.predict_on_batch() to prevent repeated TensorFlow graph re-tracing overhead on single image calls."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf9_l.add_paragraph()
    p3.text = "• Real-Time Metrics: Displays categorical prediction badge (🔴 Damaged / 🟢 Undamaged), model confidence score, and dual probability progress bars."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)
    p4 = tf9_l.add_paragraph()
    p4.text = "• Execution Command:\n   streamlit run app.py"
    p4.font.size = Pt(12)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_BLUE
    p4.space_before = Pt(8)

    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Live Demonstration Protocol")
    tb9_r = s9.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.2))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    p = tf9_r.paragraphs[0]
    p.text = "1. Model Status: App verifies models/mobilenet_v2_damage_classifier.keras is loaded."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf9_r.add_paragraph()
    p2.text = "2. Test Sample 1 (Damaged Structure): Demonstrates Damaged classification with clear visual fracture and class probability breakdown."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf9_r.add_paragraph()
    p3.text = "3. Test Sample 2 (Sound Structure): Demonstrates Undamaged classification on an intact facade component."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)
    p4 = tf9_r.add_paragraph()
    p4.text = "4. Ethical Notice: Prominently renders civil engineering disclaimer that predictions are preliminary and not safety certifications."
    p4.font.size = Pt(12)
    p4.space_before = Pt(8)

    notes9 = s9.notes_slide.notes_text_frame
    notes9.text = (
        "SPEAKER SCRIPT:\n"
        "We packaged our model into a responsive Streamlit web application. "
        "The model is cached in memory, and single-image inference runs responsively on standard CPU. "
        "Users drag and drop any JPG or PNG photo, and the app displays the predicted category, confidence score, "
        "and dual probability bars. In the review session, we will demonstrate live classifications on both damaged and undamaged samples."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 10: Limitations, Ethics & Future Work
    # ─────────────────────────────────────────────────────────────
    s10 = add_base_slide("Limitations, Ethics & Future Work")
    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Engineering Limitations & Ethics")
    tb10_l = s10.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.2))
    tf10_l = tb10_l.text_frame
    tf10_l.word_wrap = True
    p = tf10_l.paragraphs[0]
    p.text = "• Coarse Binary Scope: Distinguishes only Damaged vs. Undamaged; does not grade severity (minor cosmetic cracking vs. catastrophic shear failure)."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf10_l.add_paragraph()
    p2.text = "• Resolution Bounds: 224x224 input can overlook micro-cracks. Full-resolution hierarchical tiling would be needed for microscopic inspections."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf10_l.add_paragraph()
    p3.text = "• Distribution Shift: Trained primarily on concrete and masonry components; may exhibit lower confidence on timber, steel, or composite structures."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)
    p4 = tf10_l.add_paragraph()
    p4.text = "• Ethical Imperative: AI tools must assist, not replace, certified structural engineers."
    p4.font.size = Pt(12)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_BLUE
    p4.space_before = Pt(8)

    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Future Technical Roadmap")
    tb10_r = s10.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.2))
    tf10_r = tb10_r.text_frame
    tf10_r.word_wrap = True
    p = tf10_r.paragraphs[0]
    p.text = "1. Multi-Task Learning: Expand to Phi-Net Tasks 3-6 to classify component type, damage severity, and failure mechanism simultaneously."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p2 = tf10_r.add_paragraph()
    p2.text = "2. Explainability & Attention: Integrate Grad-CAM heatmaps so civil engineers can visually verify where the model is looking before accepting predictions."
    p2.font.size = Pt(12)
    p2.space_before = Pt(8)
    p3 = tf10_r.add_paragraph()
    p3.text = "3. Mobile & Drone Deployment: Quantize the MobileNetV2 graph to INT8 TensorFlow Lite for real-time offline edge inferencing on field drones and mobile phones."
    p3.font.size = Pt(12)
    p3.space_before = Pt(8)
    p4 = tf10_r.add_paragraph()
    p4.text = "4. Classical Baseline Comparison: Benchmark against HOG feature extractors with Support Vector Machines."
    p4.font.size = Pt(12)
    p4.space_before = Pt(8)

    notes10 = s10.notes_slide.notes_text_frame
    notes10.text = (
        "SPEAKER SCRIPT:\n"
        "To conclude responsibly, we address the limitations of our prototype. "
        "A binary classifier cannot grade whether a crack is cosmetic or structural. "
        "Furthermore, our safety disclaimer makes it clear that this prototype is strictly an educational tool. "
        "Our future roadmap includes multi-task classification across Phi-Net Tasks 3 to 6, "
        "Grad-CAM visual saliency maps for explainability, and edge quantization to TensorFlow Lite for drone inspections."
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 11: Conclusion & References
    # ─────────────────────────────────────────────────────────────
    s11 = add_base_slide("Conclusion & Key References")
    add_card(s11, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.6), "Project Summary")
    tb11_top = s11.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.3), Inches(1.8))
    tf11_t = tb11_top.text_frame
    tf11_t.word_wrap = True
    p = tf11_t.paragraphs[0]
    p.text = "• Complete Reproducible Pipeline: Successfully implemented memory-efficient data loading, verified Caffe BGR preprocessing inversion, and feature-cached transfer learning."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p2 = tf11_t.add_paragraph()
    p2.text = "• Strong Empirical Benchmark: Achieved 84.25% test accuracy and 0.9289 ROC-AUC on 1,460 official test images with 85.91% Damaged class recall."
    p2.font.size = Pt(13)
    p3 = tf11_t.add_paragraph()
    p3.text = "• Submission Deliverables: Source code in private GitHub repo, 2-page PDF report (report.pdf), presentation deck, viva guide, and running Streamlit web app."
    p3.font.size = Pt(13)

    add_card(s11, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.3), "References")
    tb11_b = s11.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(11.3), Inches(1.6))
    tf11_b = tb11_b.text_frame
    tf11_b.word_wrap = True
    p = tf11_b.paragraphs[0]
    p.text = "1. Gao, Y., & Mosalam, K. M. (2018). Deep Transfer Learning for Image‐Based Structural Damage Recognition. Comp.-Aided Civ. & Inf. Eng., 33(9), 748–768."
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED
    p2 = tf11_b.add_paragraph()
    p2.text = "2. Gao, Y., & Mosalam, K. M. (2020). PEER Hub ImageNet: A Large-Scale Benchmark Data Set. Journal of Structural Engineering, 146(10), 04020198."
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_MUTED
    p3 = tf11_b.add_paragraph()
    p3.text = "3. Sandler, M., et al. (2018). MobileNetV2: Inverted Residuals and Linear Bottlenecks. CVPR 2018, 4510–4520."
    p3.font.size = Pt(11)
    p3.font.color.rgb = COLOR_MUTED
    p4 = tf11_b.add_paragraph()
    p4.text = "4. Ho, M., & Troncoso, J. (2018). Structural Damage Image Classification. Stanford CS229 Machine Learning Project Report."
    p4.font.size = Pt(11)
    p4.font.color.rgb = COLOR_MUTED

    notes11 = s11.notes_slide.notes_text_frame
    notes11.text = (
        "SPEAKER SCRIPT:\n"
        "In conclusion, we have delivered a complete, reproducible Machine Learning pipeline for structural damage recognition. "
        "Our model achieves 84.25% test accuracy and 0.9289 ROC-AUC on the official benchmark, accompanied by an interactive demo app, "
        "our 2-page report, and full codebase. Thank you for your time, and we now welcome any questions from the panel."
    )

    prs.save(str(PPTX_OUTPUT))
    print(f"[SUCCESS] Saved PowerPoint presentation to: {PPTX_OUTPUT}")


if __name__ == "__main__":
    create_deck()
